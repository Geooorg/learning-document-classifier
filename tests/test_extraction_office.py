"""Extraktion arbeitet auf Bytes, nicht auf Pfaden – prüfbar ohne Dateien auf der Platte."""

import io
import re
import zipfile
from collections.abc import Sequence
from datetime import UTC, datetime
from pathlib import Path

import docx
import openpyxl
import pymupdf
import pytest
from conftest import pfade_fuer_vorlage

from doccls.config import RAW_DIR
from doccls.extraction import extract
from doccls.extraction.pdf import heading_by_font_size
from doccls.models import Document, SegmentKind, normalize_text

MEDIA_TYPE_PDF = "application/pdf"
MEDIA_TYPE_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MEDIA_TYPE_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def dokument(pfad: Path, media_type: str) -> tuple[Document, bytes]:
    daten = pfad.read_bytes()
    return Document.create(
        source_path=pfad.name,
        media_type=media_type,
        content=daten,
        ingested_at=datetime(2026, 1, 1, tzinfo=UTC),
    ), daten


def dokument_aus_bytes(daten: bytes, dateiname: str, media_type: str) -> Document:
    return Document.create(
        source_path=dateiname,
        media_type=media_type,
        content=daten,
        ingested_at=datetime(2026, 1, 1, tzinfo=UTC),
    )


def erste(ordner: str) -> Path:
    dateien = sorted((RAW_DIR / ordner).glob(f"*.{ordner}"))
    assert dateien, f"Keine Testdaten in {RAW_DIR / ordner}"
    return dateien[0]


def test_pdf_ergibt_ein_segment_je_seite() -> None:
    doc, daten = dokument(erste("pdf"), MEDIA_TYPE_PDF)
    segmente = extract(doc, daten)
    assert segmente
    assert all(s.kind is SegmentKind.SEITE for s in segmente)
    assert [s.index for s in segmente] == list(range(len(segmente)))
    assert segmente[0].locator == "S. 1"


def test_pdf_erfasst_alle_seiten() -> None:
    """Ein Abbruch der Seitenschleife nach der ersten Seite läge gemessen an den übrigen
    Prüfungen unbemerkt: Sie fordern eine fortlaufende Indexfolge und ``S. 1`` als erste
    Fundstelle, beides bliebe auch bei einem einzigen erfassten Segment wahr. Deshalb wird
    hier zusätzlich gegen die tatsächliche Seitenzahl der Quelle geprüft.

    Welche Datei mehrseitig ist, steht nicht mehr fest: ``_page_groups`` würfelt die
    Seitenaufteilung je Variante, damit die Seitenzahl nicht zur Abkürzung wird (Konzept
    § 6.3). Gesucht wird deshalb über alle PDF-Testdateien, bis eine mehrseitige gefunden
    ist – bei 230 PDF-Dokumenten und einer Streuung, die auf keiner Vorlage zuverlässig auf
    eine Seite zusammenfällt, kann das nicht leerlaufen.
    """
    for pfad in sorted((RAW_DIR / "pdf").glob("*.pdf")):
        doc, daten = dokument(pfad, MEDIA_TYPE_PDF)
        with pymupdf.open(stream=daten, filetype="pdf") as pdf:
            seitenzahl = len(pdf)
        if seitenzahl > 1:
            segmente = extract(doc, daten)
            assert len(segmente) == seitenzahl
            return
    raise AssertionError("Keine mehrseitige PDF-Testdatei gefunden – Prüfung liefe leer")


def test_pdf_segmenttext_deckt_gesamten_seitentext_ab() -> None:
    """Eine Kürzung wie ``normalize_text(...)[:60]`` änderte weder Segmentzahl noch
    Indexfolge noch den ersten Locator – alles, was die übrigen Tests prüfen. Deshalb wird
    hier die Zeichenzahl der Quelle (unabhängig, aber mit derselben Normalisierung wie die
    Extraktion, gemessen) gegen die Summe über alle Segmenttexte geprüft.
    """
    pfad = erste("pdf")
    doc, daten = dokument(pfad, MEDIA_TYPE_PDF)
    segmente = extract(doc, daten)
    with pymupdf.open(stream=daten, filetype="pdf") as pdf:
        erwartete_laenge = sum(
            len(text) for seite in pdf if (text := normalize_text(seite.get_text("text")))
        )
    assert erwartete_laenge > 60, "Testdatei liefert zu wenig Text – Prüfung liefe leer"
    assert sum(len(s.text) for s in segmente) == erwartete_laenge


def test_heading_by_font_size_erkennt_deutlich_groessere_schrift() -> None:
    """``MIN_HEADING_SIZE_RATIO`` und die Median-Berechnung des Fließtexts haben bisher
    keinen einzigen Test: Eine Stilllegung (z. B. ein unerreichbar hoher Schwellwert)
    ließe jede Überschrift zu ``None`` werden, ohne dass eine Prüfung auf Segmentzahl oder
    -reihenfolge das bemerken würde.
    """
    with pymupdf.open() as pdf:
        seite = pdf.new_page()
        seite.insert_text((72, 100), "Titelzeile Rechnung", fontsize=24)
        for i in range(5):
            seite.insert_text((72, 150 + i * 15), f"Fliesstext Zeile {i} mit Inhalt.", fontsize=10)
        ergebnis = heading_by_font_size(seite)
    assert ergebnis == "Titelzeile Rechnung"


def test_heading_by_font_size_ignoriert_leichten_groessenunterschied() -> None:
    """Der Schwellwert greift erst ab einem deutlichen Unterschied zum (über den Median
    bestimmten) Fließtext – knapp darüber reicht nicht, sonst würde jede zufällig etwas
    größere Zeile fälschlich zur Überschrift.
    """
    with pymupdf.open() as pdf:
        seite = pdf.new_page()
        seite.insert_text((72, 100), "Fast wie Fliesstext", fontsize=11.3)
        for i in range(5):
            seite.insert_text((72, 150 + i * 15), f"Fliesstext Zeile {i} mit Inhalt.", fontsize=10)
        ergebnis = heading_by_font_size(seite)
    assert ergebnis is None


def test_heading_by_font_size_ignoriert_lange_zeilen() -> None:
    """``MAX_HEADING_LENGTH`` schließt Fließtext aus, der zufällig groß gesetzt ist (etwa
    ein hervorgehobener Absatz) – lang und groß ist kein Widerspruch, aber keine
    Überschrift.
    """
    with pymupdf.open() as pdf:
        seite = pdf.new_page(width=2000, height=400)
        lang = (
            "Eine sehr lange Ueberschriftenzeile die eindeutig laenger ist als achtzig "
            "Zeichen insgesamt und auf einer Zeile bleibt"
        )
        seite.insert_text((72, 100), lang, fontsize=24)
        for i in range(5):
            seite.insert_text((72, 150 + i * 15), f"Fliesstext Zeile {i} mit Inhalt.", fontsize=10)
        ergebnis = heading_by_font_size(seite)
    assert ergebnis is None


def test_docx_liefert_abschnitte_mit_ueberschriften() -> None:
    doc, daten = dokument(erste("docx"), MEDIA_TYPE_DOCX)
    segmente = extract(doc, daten)
    assert any(s.heading for s in segmente)
    assert all(s.kind is SegmentKind.ABSCHNITT for s in segmente)


def test_docx_segmenttext_deckt_alle_absaetze_ab() -> None:
    """``absaetze[:] = [text]`` ließe je Abschnitt nur den letzten Absatz übrig –
    Segmentzahl, Überschriften und Segmentart blieben unverändert, nur der Inhalt
    schrumpft. Geprüft wird deshalb die Zeichenzahl der einzeln gezählten, nicht aus
    Überschriften stammenden Absätze gegen die Summe der (nicht aus Tabellen stammenden)
    Abschnittstexte: Da ``normalize_text`` die Absätze eines Abschnitts nur durch je ein
    Leerzeichen verbindet, kann die Summe der Segmenttexte nie kleiner sein als die Summe
    der Absätze – außer es geht unterwegs einer verloren.
    """
    pfad = erste("docx")
    doc, daten = dokument(pfad, MEDIA_TYPE_DOCX)
    segmente = extract(doc, daten)

    quelle = docx.Document(io.BytesIO(daten))
    absatz_texte = []
    for absatz in quelle.paragraphs:
        text = absatz.text.strip()
        if not text:
            continue
        stil = absatz.style.name if absatz.style is not None else None
        if stil is not None and stil.startswith("Heading"):
            continue
        absatz_texte.append(text)

    abschnitts_segmente = [s for s in segmente if not s.locator.startswith("Tabelle")]
    assert len(absatz_texte) > len(abschnitts_segmente), (
        "Testdatei bietet keine Mehrfach-Absatz-Abschnitte – Prüfung liefe leer"
    )
    erwartete_mindestlaenge = sum(len(text) for text in absatz_texte)
    assert sum(len(s.text) for s in abschnitts_segmente) >= erwartete_mindestlaenge


def test_docx_tabellen_ergeben_je_ein_segment() -> None:
    """Eine leer laufende Tabellenschleife fiele nicht durch die Segmentzahl allein auf,
    solange das Dokument auch Absatzabschnitte enthält. Geprüft wird deshalb die Anzahl
    der Tabellen in der Quelle gegen die Anzahl der Tabellensegmente.
    """
    pfad = pfade_fuer_vorlage("GUTSCHRIFT-bonus")[0]
    doc, daten = dokument(pfad, MEDIA_TYPE_DOCX)

    quelle = docx.Document(io.BytesIO(daten))
    assert quelle.tables, "Testdatei enthält keine Tabelle – Prüfung liefe leer"

    segmente = extract(doc, daten)
    tabellen_segmente = [s for s in segmente if s.locator.startswith("Tabelle")]
    assert len(tabellen_segmente) == len(quelle.tables)


def test_docx_tabelle_ohne_text_erzeugt_kein_segment() -> None:
    """``" | ".join`` über leere Zellen ergibt immer Trennzeichen (``"| | | |"`` ist nicht
    leer) – der Wächter ``if text:`` konnte deshalb nie greifen. Layout-Tabellen ohne
    Inhalt sind in echten Word-Briefen der Normalfall.
    """
    quelle = docx.Document()
    quelle.add_paragraph("Sehr geehrte Damen und Herren,")
    quelle.add_table(rows=2, cols=3)
    puffer = io.BytesIO()
    quelle.save(puffer)
    daten = puffer.getvalue()
    doc = dokument_aus_bytes(daten, "brief.docx", MEDIA_TYPE_DOCX)

    segmente = extract(doc, daten)
    assert not any(s.locator.startswith("Tabelle") for s in segmente)


def test_docx_ueberschrift_ohne_folgeabsatz_geht_nicht_verloren() -> None:
    """K2 Nachweis a: Ein Dokumenttitel und eine Zwischenüberschrift, denen jeweils keine
    eigenen Absätze folgen (der Zwischenüberschrift folgt direkt eine Tabelle), wurden
    bisher stillschweigend verworfen – ``abschliessen`` schrieb nur bei nichtleerem
    ``absaetze``. Zusätzlich muss die Tabelle die zuletzt gesehene Überschrift tragen statt
    ``None``, sonst ist der Bezug zwischen Abschnitt und Tabelle gekappt.
    """
    quelle = docx.Document()
    quelle.add_heading("Rechnung RE-2026-1234", level=1)
    quelle.add_heading("Positionen", level=2)
    tabelle = quelle.add_table(rows=2, cols=2)
    tabelle.rows[0].cells[0].text = "Leistung"
    tabelle.rows[0].cells[1].text = "Betrag"
    tabelle.rows[1].cells[0].text = "Wartung"
    tabelle.rows[1].cells[1].text = "750,00 EUR"
    quelle.add_heading("Zahlung", level=2)
    quelle.add_paragraph("Zahlbar bis 30.04.2026 ohne Abzug.")
    puffer = io.BytesIO()
    quelle.save(puffer)
    daten = puffer.getvalue()
    doc = dokument_aus_bytes(daten, "rechnung.docx", MEDIA_TYPE_DOCX)

    segmente = extract(doc, daten)
    ueberschriften = {s.heading for s in segmente if s.heading}
    assert "Rechnung RE-2026-1234" in ueberschriften, "Dokumenttitel ohne Folgeabsatz verloren"
    assert "Positionen" in ueberschriften, "Zwischenüberschrift ohne Folgeabsatz verloren"

    tabellen_segment = next(s for s in segmente if s.locator == "Tabelle 1")
    assert tabellen_segment.heading == "Positionen", (
        "Der Bezug zwischen Überschrift und Tabelle ist gekappt"
    )
    assert "750,00" in tabellen_segment.text


def test_docx_verschachtelte_tabelle_wird_gelesen() -> None:
    """K2 Nachweis b: ``zelle.text`` sieht keine in eine Zelle eingebettete Tabelle. Eine
    Rahmentabelle, deren einzige Zelle den eigentlichen Inhalt in einer verschachtelten
    Tabelle trägt, lieferte bisher null Segmente – ein vollständiger, stiller Inhaltsverlust,
    mit dem das Dokument dennoch als erfolgreich eingelesen gegolten hätte.
    """
    quelle = docx.Document()
    quelle.add_heading("Rechnung", level=1)
    aussen = quelle.add_table(rows=1, cols=1)
    innen = aussen.rows[0].cells[0].add_table(rows=1, cols=2)
    innen.rows[0].cells[0].text = "Summe"
    innen.rows[0].cells[1].text = "12.000,00 EUR"
    puffer = io.BytesIO()
    quelle.save(puffer)
    daten = puffer.getvalue()
    doc = dokument_aus_bytes(daten, "verschachtelt.docx", MEDIA_TYPE_DOCX)

    segmente = extract(doc, daten)
    assert segmente, "Verschachtelte Tabelle darf nicht zu einem inhaltslosen Dokument führen"
    assert any("12.000,00" in s.text for s in segmente), "Inhalt der verschachtelten Tabelle fehlt"


def test_xlsx_liefert_ein_segment_je_blatt() -> None:
    doc, daten = dokument(erste("xlsx"), MEDIA_TYPE_XLSX)
    segmente = extract(doc, daten)
    assert segmente
    assert all(s.kind is SegmentKind.BLATT for s in segmente)


def test_tabellenzeile_bleibt_als_einheit_beisammen() -> None:
    """Zerfiele eine Zeile in Zellen, verlöre man den Bezug zwischen Leistung und Betrag –
    und genau dieser Bezug unterscheidet eine Rechnung von einer Liste.

    Geprüft wird deshalb, dass Spaltenname und Wert im **selben** Segment stehen. Eine
    Suche nach "EUR" im zusammengefügten Gesamttext wäre auch dann wahr, wenn jede Zelle
    ein eigenes Segment bekäme – also genau im Fehlerfall.
    """
    pfad = pfade_fuer_vorlage("RECHNUNG-tabelle")[0]
    doc, daten = dokument(pfad, MEDIA_TYPE_XLSX)
    treffer = [s for s in extract(doc, daten) if "Einzelpreis:" in s.text]
    assert treffer, "Kein Segment enthält die Spalte 'Einzelpreis'"
    zeile = treffer[0].text
    assert "Leistung:" in zeile, "Spalten derselben Zeile liegen in verschiedenen Segmenten"
    assert re.search(r"Einzelpreis: [\d.,]+ EUR", zeile), "Spaltenname und Wert sind getrennt"


def test_xlsx_segmenttext_erfasst_alle_datenzeilen() -> None:
    """``rest = zeilen[1:2]`` ließe je Blatt nur eine Datenzeile durch – Segmentzahl und
    Blattname blieben unverändert. Geprüft wird deshalb, wie oft eine in jeder Zeile
    garantiert gefüllte Spalte im Segmenttext auftaucht, gegen die tatsächliche Anzahl der
    Datenzeilen (Kopfzeile ausgenommen), unabhängig ermittelt.
    """
    pfad = pfade_fuer_vorlage("RECHNUNG-tabelle")[0]
    doc, daten = dokument(pfad, MEDIA_TYPE_XLSX)

    mappe = openpyxl.load_workbook(io.BytesIO(daten), read_only=True, data_only=True)
    try:
        blatt_titel: str | None = None
        alle_zeilen: Sequence[tuple[object, ...]] = ()
        for blatt in mappe.worksheets:
            blatt.reset_dimensions()
            zeilen = list(blatt.iter_rows(values_only=True))
            if zeilen and "Einzelpreis" in zeilen[0]:
                blatt_titel, alle_zeilen = blatt.title, zeilen
                break
    finally:
        mappe.close()
    assert blatt_titel is not None, "Kein Blatt mit Spalte 'Einzelpreis' gefunden"
    datenzeilen = sum(
        1 for zeile in alle_zeilen[1:] if any(zelle not in (None, "") for zelle in zeile)
    )
    assert datenzeilen > 1, "Testdatei bietet nur eine Datenzeile – Prüfung liefe leer"

    segmente = extract(doc, daten)
    text = next(s.text for s in segmente if s.locator == f"Blatt {blatt_titel}")
    assert text.count("Einzelpreis:") == datenzeilen


def _xlsx_mit_kaputter_dimension() -> bytes:
    """Baut eine Mappe mit Kopf- plus drei Positionszeilen und ersetzt danach das
    ``<dimension ref="…">`` im Blatt-XML durch den zu kleinen Platzhalter ``A1`` – ganz
    ohne Artefakt auf der Platte. ``<dimension ref="A1"/>`` ist der übliche Platzhalter
    von Erzeugern (ERP- und CSV-Exporter, Streaming-Writer), die die Ausdehnung beim
    Schreiben nicht vorab kennen; Excel korrigiert das beim Öffnen, openpyxl
    (``read_only=True``) nicht.
    """
    mappe = openpyxl.Workbook()
    blatt = mappe.active
    assert blatt is not None
    blatt.title = "Positionen"
    blatt.append(["Leistung", "Menge", "Einzelpreis", "Betrag"])
    for i in range(3):
        blatt.append([f"Abnahme {i}", 1, "900,00", "900,00"])
    puffer = io.BytesIO()
    mappe.save(puffer)

    quelle = zipfile.ZipFile(io.BytesIO(puffer.getvalue()))
    ziel_puffer = io.BytesIO()
    with zipfile.ZipFile(ziel_puffer, "w", zipfile.ZIP_DEFLATED) as ziel:
        for eintrag in quelle.infolist():
            inhalt = quelle.read(eintrag.filename)
            if eintrag.filename == "xl/worksheets/sheet1.xml":
                inhalt = re.sub(rb'<dimension ref="[^"]*"/>', b'<dimension ref="A1"/>', inhalt)
            ziel.writestr(eintrag, inhalt)
    return ziel_puffer.getvalue()


def test_xlsx_ueberlebt_falsche_dimension_angabe() -> None:
    """``read_only=True`` vertraut dem ``<dimension>``-Attribut der Blatt-XML, statt zu
    messen, was tatsächlich da ist. Steht dort ein zu kleiner Platzhalter, verwirft
    openpyxl Spalten und Zeilen kommentarlos – hier müssten sonst alle drei Positionen
    mit ihren Beträgen ankommen, nicht nur ein Bruchteil.
    """
    daten = _xlsx_mit_kaputter_dimension()
    doc = dokument_aus_bytes(daten, "rechnung.xlsx", MEDIA_TYPE_XLSX)

    segmente = extract(doc, daten)
    assert len(segmente) == 1
    text = segmente[0].text
    assert "900,00" in text
    assert text.count("Einzelpreis:") == 3, "Nur ein Teil der Datenzeilen ist angekommen"
    assert text.count("Abnahme") == 3, "Nur ein Teil der Datenzeilen ist angekommen"


def _xlsx_mit_kurzer_kopfzeile() -> bytes:
    """Kopfzeile mit zwei gefüllten Zellen, Datenzeile mit vier – openpyxl polstert die
    Kopfzeile beim Lesen auf die Breite der breitesten Zeile auf."""
    mappe = openpyxl.Workbook()
    blatt = mappe.active
    assert blatt is not None
    blatt.title = "Test"
    blatt["A1"] = "Leistung"
    blatt["B1"] = "Menge"
    blatt["A2"] = "Abnahme"
    blatt["B2"] = 1
    blatt["C2"] = 10
    blatt["D2"] = 10
    puffer = io.BytesIO()
    mappe.save(puffer)
    return puffer.getvalue()


def test_xlsx_kopfzeile_ohne_leere_spalten_im_text() -> None:
    """openpyxl polstert jede gelesene Zeile auf die Spaltenzahl des Blattes auf. Hat die
    Kopfzeile selbst weniger gefüllte Zellen als eine spätere Datenzeile, erscheinen die
    aufgefüllten Leerzellen als zusätzliche Trennzeichen (``"| |"``) in der dargestellten
    Kopfzeile. Die Zuordnung Spaltenname → Wert in der Datenzeile darf davon unberührt
    bleiben.
    """
    daten = _xlsx_mit_kurzer_kopfzeile()
    doc = dokument_aus_bytes(daten, "kurz.xlsx", MEDIA_TYPE_XLSX)

    segmente = extract(doc, daten)
    assert len(segmente) == 1
    assert "| |" not in segmente[0].text
    assert segmente[0].text == "Leistung | Menge | Leistung: Abnahme, Menge: 1, : 10, : 10"


def _xlsx_mit_summenzelle(mit_formel: bool) -> bytes:
    """Identische Mappe, einmal mit einer Formel in der letzten Spalte, einmal mit dem
    fertigen Wert – wie im Schlussbefund (K3) beschrieben. Eine per API mit ``blatt.append``
    gesetzte Formel hat keinen zwischengespeicherten Wert, genau wie bei ERP- und
    Streaming-Exportern, die den Cache beim Schreiben nicht mitführen.
    """
    mappe = openpyxl.Workbook()
    blatt = mappe.active
    assert blatt is not None
    blatt.title = "Rechnung"
    blatt.append(["Pos", "Menge", "Einzelpreis", "Summe"])
    letzte_spalte: str | int = "=B2*C2" if mit_formel else 750
    blatt.append([1, 3, 250, letzte_spalte])
    puffer = io.BytesIO()
    mappe.save(puffer)
    return puffer.getvalue()


def test_xlsx_formelzelle_ohne_cache_wird_nicht_verschluckt() -> None:
    """K3: ``data_only=True`` liefert für eine Formelzelle ohne zwischengespeicherten Wert
    ``None`` – ununterscheidbar von einer echten Leerzelle. Die Spaltenüberschrift „Summe"
    bliebe stehen, der Betrag würde aber lautlos verschwinden, obwohl gerade er das
    entscheidende Unterscheidungsmerkmal zwischen Rechnung und Gutschrift ist.
    """
    daten = _xlsx_mit_summenzelle(mit_formel=True)
    doc = dokument_aus_bytes(daten, "formel.xlsx", MEDIA_TYPE_XLSX)

    segmente = extract(doc, daten)
    assert len(segmente) == 1
    text = segmente[0].text
    assert "Summe:" in text, "Spalte 'Summe' fehlt vollständig"
    nach_summe = text.split("Summe:", 1)[1].strip()
    assert nach_summe and not nach_summe.startswith(("None", ",")), (
        f"Formelzelle wurde verschluckt statt sichtbar gemacht: {text!r}"
    )
    assert "=B2*C2" in text, "Die Formel selbst sollte als Ersatz für den fehlenden Wert stehen"


def test_xlsx_formelzelle_mit_cache_liefert_den_wert() -> None:
    """Gegenprobe zu K3: Trägt die Zelle bereits den fertigen Wert (wie eine von Excel
    gespeicherte Datei mit aufgelöstem Formel-Cache), muss dieser unverändert ankommen.
    """
    daten = _xlsx_mit_summenzelle(mit_formel=False)
    doc = dokument_aus_bytes(daten, "formel_mit_wert.xlsx", MEDIA_TYPE_XLSX)

    segmente = extract(doc, daten)
    assert len(segmente) == 1
    assert "Summe: 750" in segmente[0].text


def test_xlsx_ueberlebt_falsche_dimension_angabe_trotz_zweitem_ladevorgang() -> None:
    """Der K3-Fix lädt die Mappe zweimal (``data_only=True`` und ``False``). Beide
    Ladevorgänge müssen ``reset_dimensions()`` unabhängig anwenden – sonst gälte die
    Korrektur aus einem früheren Befund (falsches ``<dimension ref>``) nur noch für den
    ersten Durchlauf und der zweite würde Zeilen unterschiedlicher Länge liefern.
    """
    daten = _xlsx_mit_kaputter_dimension()
    doc = dokument_aus_bytes(daten, "rechnung.xlsx", MEDIA_TYPE_XLSX)

    segmente = extract(doc, daten)
    assert len(segmente) == 1
    text = segmente[0].text
    assert text.count("Einzelpreis:") == 3
    assert text.count("Abnahme") == 3


def test_text_ist_normalisiert() -> None:
    doc, daten = dokument(erste("pdf"), MEDIA_TYPE_PDF)
    segmente = extract(doc, daten)
    assert segmente, "Ohne Segmente prüft die Schleife nichts"
    for segment in segmente:
        assert "  " not in segment.text
        assert "\n" not in segment.text


def test_unbekanntes_format_wird_abgelehnt() -> None:
    doc, _ = dokument(erste("pdf"), MEDIA_TYPE_PDF)
    with pytest.raises(ValueError, match="Nicht unterstützt"):
        extract(doc, b"\x00\x01\x02 irgendwas")


# Aufgabe 3 verlangt deterministische index-Vergabe: derselbe Bytes-Input muss bei jeder
# Extraktion dieselbe Reihenfolge und Indexfolge ergeben. Bibliotheken wie openpyxl
# garantieren das nicht von sich aus für jede Konstellation – deshalb wird hier gemessen,
# nicht angenommen.
DATEIEN_JE_FORMAT: dict[str, str] = {
    "pdf": MEDIA_TYPE_PDF,
    "docx": MEDIA_TYPE_DOCX,
    "xlsx": MEDIA_TYPE_XLSX,
}


@pytest.mark.parametrize("ordner,media_type", DATEIEN_JE_FORMAT.items())
def test_extraktion_ist_deterministisch(ordner: str, media_type: str) -> None:
    doc, daten = dokument(erste(ordner), media_type)
    erster_lauf = extract(doc, daten)
    zweiter_lauf = extract(doc, daten)
    assert erster_lauf, "Ohne Segmente prüft der Vergleich nichts"
    assert erster_lauf == zweiter_lauf
    assert [s.index for s in erster_lauf] == [s.index for s in zweiter_lauf]
